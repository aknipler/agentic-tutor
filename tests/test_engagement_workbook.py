"""Tests for utils.engagement_workbook, using synthetic records and no database.

The workbook's score and rank cells are formulas, and openpyxl doesn't calculate formulas. These
tests check the layout the formulas depend on and the values the script writes. The formulas
themselves are checked by opening a workbook in a spreadsheet program (see the validation
section of docs/engagement-report-plan.md).

Run them from the repository root with

    .venv/Scripts/python.exe -m unittest discover -s tests -t .
"""
import unittest
from datetime import date, datetime, timedelta

from openpyxl.utils import get_column_letter

from utils.engagement import PRIZE_1_MEASURES, PRIZE_2_MEASURES, Settings, StudentRecords, build_report
from utils.engagement_checks import CheckResult
from utils.engagement_workbook import (
    FIRST_DATA_ROW,
    HEADER_ROW,
    THRESHOLD_ROWS,
    TIMING_FLAG,
    WEIGHT_ROW,
    RunInfo,
    build_workbook,
)

REPORT_DATE = date(2026, 10, 17)


def utc(month, day, hour=0, minute=0):
    return datetime(2026, month, day, hour, minute)


def submission(when, level=2, question="Q1"):
    return {"module": "1", "question": question, "grade": level / 2, "timestamp": when}


def conversation(pairs, topic="Topic A"):
    return {"module": "1", "topic": topic, "conversation": pairs, "timestamp": utc(8, 3)}


def cohort():
    weekly = [submission(utc(8, 3) + timedelta(days=7 * week), level=week % 3) for week in range(5)]
    return [
        StudentRecords("both", [conversation([["why?", "reply", utc(8, 4, 1)], ["ok", "reply"]])], weekly),
        StudentRecords("tutor_only", [conversation([[f"message {i}", "reply"] for i in range(12)])], []),
        StudentRecords("assessor_only", [], [submission(utc(9, 1), 1), submission(utc(9, 2), 2)]),
        StudentRecords("=SUM(1,2)", [], [submission(utc(9, 5), 1, question="=1+1")]),
    ]


def build(students=None, run=None, settings=None):
    report = build_report(cohort() if students is None else students, REPORT_DATE, settings or Settings())
    run = run or RunInfo("17 Oct 2026 10:00 AEDT", "abc1234")
    return report, build_workbook(report, run)


def column_of(ws, header, row=HEADER_ROW):
    for cell in ws[row]:
        if cell.value == header:
            return cell.column
    raise AssertionError(f"no column headed {header!r} on {ws.title}")


def rows_by_student(ws, first_row=FIRST_DATA_ROW):
    return {row[0].value: row for row in ws.iter_rows(min_row=first_row) if row[0].value is not None}


class LayoutTests(unittest.TestCase):
    def test_the_sheets_come_in_the_planned_order(self):
        _, wb = build()

        self.assertEqual(wb.sheetnames, ["About", "Students", "Weekly", "Attempts", "Prize 1", "Prize 2"])

    def test_an_empty_report_still_builds(self):
        _, wb = build(students=[])

        self.assertEqual(wb.sheetnames[-1], "Prize 2")
        self.assertEqual(wb["Prize 1"].tables, {})

    def test_text_that_looks_like_a_formula_is_stored_as_text(self):
        _, wb = build()

        for sheet in ("Students", "Attempts", "Prize 1", "Prize 2"):
            first = wb[sheet].cell(row=FIRST_DATA_ROW if "Prize" in sheet else 2, column=1)
            self.assertEqual(first.data_type, "s", sheet)
        row = rows_by_student(wb["Prize 1"])["=SUM(1,2)"]
        self.assertEqual(row[0].data_type, "s")
        question = [r for r in wb["Attempts"].iter_rows(min_row=2) if r[0].value == "=SUM(1,2)"][0][2]
        self.assertEqual((question.value, question.data_type), ("=1+1", "s"))

    def test_table_headings_are_unique(self):
        _, wb = build()

        for ws in wb.worksheets:
            for table in ws.tables.values():
                first, last = table.ref.split(":")
                row = ws[first.lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ")]
                headings = [cell.value for cell in row]
                headings = [h for h in headings if h is not None]
                self.assertEqual(len(headings), len(set(headings)), ws.title)

    def test_about_text_follows_the_house_style(self):
        _, wb = build(run=RunInfo("now", "abc", [CheckResult("A check", True, "All matched.")], ["Done by hand."]))

        for row in wb["About"].iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    self.assertNotIn("—", cell.value)
                    self.assertNotIn(";", cell.value)


class StudentsSheetTests(unittest.TestCase):
    def test_one_row_per_student_with_interactions(self):
        report, wb = build(students=cohort() + [StudentRecords("silent", [], [])])
        ws = wb["Students"]

        listed = [row[0].value for row in ws.iter_rows(min_row=2)]

        self.assertEqual(listed, [s.user_id for s in report.students])
        self.assertNotIn("silent", listed)

    def test_counts_and_flag_match_the_metrics(self):
        report, wb = build()
        ws = wb["Students"]
        by_id = {s.user_id: s for s in report.students}
        rows = rows_by_student(ws, first_row=2)

        tutor = column_of(ws, "Tutor messages", row=1)
        flag = column_of(ws, "Tutor timing", row=1)
        self.assertEqual(rows["tutor_only"][tutor - 1].value, by_id["tutor_only"].tutor_messages)
        self.assertEqual(rows["tutor_only"][flag - 1].value, TIMING_FLAG)
        self.assertIsNone(rows["both"][flag - 1].value)

    def test_a_missing_figure_is_blank_not_zero(self):
        _, wb = build()
        ws = wb["Students"]

        rows = rows_by_student(ws, first_row=2)

        self.assertIsNone(rows["tutor_only"][column_of(ws, "Attempts per question", row=1) - 1].value)


class WeeklySheetTests(unittest.TestCase):
    def test_each_students_week_cells_add_up_to_their_totals(self):
        report, wb = build()
        ws = wb["Weekly"]
        weeks = len(report.weeks)
        rows = rows_by_student(ws, first_row=3)

        for student in report.students:
            cells = [cell.value or 0 for cell in rows[student.user_id]]
            tutor, undated, assessor = cells[1:1 + weeks], cells[1 + weeks], cells[2 + weeks:2 + 2 * weeks]
            self.assertEqual(sum(tutor) + undated, student.tutor_messages, student.user_id)
            self.assertEqual(sum(assessor), student.assessor_submissions, student.user_id)

    def test_undated_messages_have_their_own_column(self):
        report, wb = build()
        ws = wb["Weekly"]
        rows = rows_by_student(ws, first_row=3)

        undated = column_of(ws, "Undated", row=2)

        self.assertEqual(rows["tutor_only"][undated - 1].value, 12)
        self.assertIsNone(rows["assessor_only"][undated - 1].value)

    def test_week_headings_are_mondays(self):
        report, wb = build()

        headings = [cell.value for cell in wb["Weekly"][2][1:1 + len(report.weeks)]]

        self.assertEqual(headings, report.weeks)
        self.assertTrue(all(week.weekday() == 0 for week in headings))


class AttemptsSheetTests(unittest.TestCase):
    def test_one_row_per_student_and_question(self):
        report, wb = build()
        ws = wb["Attempts"]

        self.assertEqual(ws.max_row - 1, len(report.attempts))
        both = [r for r in ws.iter_rows(min_row=2) if r[0].value == "both"]
        self.assertEqual(len(both), 1)
        self.assertEqual(both[0][3].value, 5)
        self.assertEqual(both[0][4].value, "0 → 1 → 2 → 0 → 1")
        self.assertEqual((both[0][5].value, both[0][6].value), (0, 2))


class PrizeSheetTests(unittest.TestCase):
    def percentile_columns(self, ws, measures):
        return [column_of(ws, f"{m.label} percentile") for m in measures]

    def test_weights_sit_directly_above_contiguous_percentile_columns(self):
        report, wb = build()

        for sheet, table in (("Prize 1", report.prize_1), ("Prize 2", report.prize_2)):
            ws = wb[sheet]
            columns = self.percentile_columns(ws, table.measures)
            self.assertEqual(columns, list(range(columns[0], columns[0] + len(columns))), sheet)
            for column, measure in zip(columns, table.measures):
                self.assertEqual(ws.cell(row=WEIGHT_ROW, column=column).value, measure.weight, f"{sheet} {measure.key}")

    def test_score_and_rank_formulas_read_the_weights_row_and_their_own_row(self):
        report, wb = build()
        ws = wb["Prize 1"]
        columns = self.percentile_columns(ws, PRIZE_1_MEASURES)
        first, last = get_column_letter(columns[0]), get_column_letter(columns[-1])
        score = ws.cell(row=FIRST_DATA_ROW, column=column_of(ws, "Score")).value
        rank = ws.cell(row=FIRST_DATA_ROW, column=column_of(ws, "Rank")).value
        count = len(report.students)

        self.assertEqual(
            score,
            f'=IFERROR(ROUND(SUMPRODUCT(${first}${WEIGHT_ROW}:${last}${WEIGHT_ROW},{first}{FIRST_DATA_ROW}:{last}{FIRST_DATA_ROW})'
            f'/SUMPRODUCT(${first}${WEIGHT_ROW}:${last}${WEIGHT_ROW},--ISNUMBER({first}{FIRST_DATA_ROW}:{last}{FIRST_DATA_ROW})),6),"")',
        )
        self.assertIn(f"$C${FIRST_DATA_ROW}:$C${FIRST_DATA_ROW + count - 1}", rank)
        self.assertIn(f">C{FIRST_DATA_ROW}", rank)

    def test_every_row_has_its_own_formulas(self):
        report, wb = build()
        ws = wb["Prize 2"]

        for row in range(FIRST_DATA_ROW, FIRST_DATA_ROW + len(report.students)):
            for header in ("Score", "Rank", "Eligible"):
                value = ws.cell(row=row, column=column_of(ws, header)).value
                self.assertTrue(value.startswith("="), f"{header} row {row}")
                self.assertIn(str(row), value)

    def test_eligibility_reads_the_threshold_cells(self):
        settings = Settings(min_weeks=4, min_interactions=7)
        report, wb = build(settings=settings)
        ws = wb["Prize 2"]
        eligible = column_of(ws, "Eligible")
        weeks = get_column_letter(column_of(ws, "Active weeks"))
        interactions = get_column_letter(column_of(ws, "Interactions"))

        self.assertEqual(ws.cell(row=THRESHOLD_ROWS[0], column=eligible).value, 4)
        self.assertEqual(ws.cell(row=THRESHOLD_ROWS[1], column=eligible).value, 7)
        formula = ws.cell(row=FIRST_DATA_ROW, column=eligible).value
        letter = get_column_letter(eligible)
        self.assertEqual(
            formula,
            f'=IF(AND({weeks}{FIRST_DATA_ROW}>=${letter}${THRESHOLD_ROWS[0]},'
            f'{interactions}{FIRST_DATA_ROW}>=${letter}${THRESHOLD_ROWS[1]}),"Yes","No")',
        )

    def test_students_with_no_data_for_a_measure_get_a_blank_percentile_and_raw_value(self):
        _, wb = build()
        ws = wb["Prize 1"]
        rows = rows_by_student(ws)

        for header in ("Tutor messages percentile", "Tutor messages (raw)"):
            self.assertIsNone(rows["assessor_only"][column_of(ws, header) - 1].value, header)
        self.assertEqual(rows["tutor_only"][column_of(ws, "Tutor messages (raw)") - 1].value, 12)

    def test_rates_show_counts_raw_rate_and_adjusted_rate(self):
        report, wb = build()
        ws = wb["Prize 2"]
        rows = rows_by_student(ws)
        both = next(s for s in report.students if s.user_id == "both")

        successes = column_of(ws, "Improvement after feedback (questions improved)")
        opportunities = column_of(ws, "Improvement after feedback (questions first scored below 2)")
        raw = column_of(ws, "Improvement after feedback (raw rate)")
        adjusted = column_of(ws, "Improvement after feedback (adjusted rate)")

        self.assertEqual(rows["both"][successes - 1].value, both.questions_improved)
        self.assertEqual(rows["both"][opportunities - 1].value, both.questions_first_below_2)
        self.assertAlmostEqual(rows["both"][raw - 1].value, both.improved_share)
        self.assertIsNotNone(rows["both"][adjusted - 1].value)
        # A student with no question first scored below 2 has no rate at all.
        self.assertIsNone(rows["tutor_only"][raw - 1].value)

    def test_the_flag_column_marks_incomplete_tutor_timing(self):
        _, wb = build()

        for sheet in ("Prize 1", "Prize 2"):
            ws = wb[sheet]
            rows = rows_by_student(ws)
            flag = column_of(ws, "Tutor timing") - 1
            self.assertEqual(rows["tutor_only"][flag].value, TIMING_FLAG, sheet)
            self.assertIsNone(rows["both"][flag].value, sheet)

    def test_prize_2_lists_every_measure_with_its_starting_weight(self):
        _, wb = build()
        ws = wb["Prize 2"]

        weights = {m.label: ws.cell(row=WEIGHT_ROW, column=column_of(ws, f"{m.label} percentile")).value for m in PRIZE_2_MEASURES}

        self.assertEqual(weights["Improvement after feedback"], 1)
        self.assertEqual(weights["Revisiting"], 1)
        self.assertEqual(weights["Follow-up questions"], 1)
        self.assertEqual(sum(weights.values()), 3)


class AboutSheetTests(unittest.TestCase):
    def test_the_run_and_validation_results_are_recorded(self):
        checks = [CheckResult("Totals match", True, "All 4 students match."), CheckResult("Attempts match", False, "One differs.", ("s1: module 1 question 1, log 2, progress 1",))]
        _, wb = build(run=RunInfo("17 Oct 2026 10:00 AEDT", "abc1234", checks, ["Traced two students by hand."]))

        text = "\n".join(str(cell.value) for row in wb["About"].iter_rows() for cell in row if cell.value is not None)

        for expected in ("17 Oct 2026 10:00 AEDT", "abc1234", "Totals match. Passed", "Attempts match. Failed", "s1: module 1 question 1, log 2, progress 1", "Traced two students by hand."):
            self.assertIn(expected, text)

    def test_settings_are_recorded(self):
        _, wb = build(settings=Settings(session_gap=timedelta(minutes=45), shrink_k=8, min_weeks=2, min_interactions=6))
        ws = wb["About"]

        values = {row[0].value: row[1].value for row in ws.iter_rows() if row[0].value and len(row) > 1}

        self.assertEqual(values["Session gap, minutes"], 45)
        self.assertEqual(values["Shrinkage constant k"], 8)
        self.assertEqual(values["Minimum active weeks, starting value"], 2)
        self.assertEqual(values["Minimum interactions, starting value"], 6)


if __name__ == "__main__":
    unittest.main()
