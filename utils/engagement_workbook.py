"""Write the engagement report as an Excel workbook.

The workbook has six sheets (About, Students, Weekly, Attempts, Prize 1 and Prize 2). The
Students, Weekly and Attempts sheets hold values worked out by utils/engagement.py. The prize
sheets hold raw values and percentiles as values, and write the score, rank and eligibility
cells as formulas, so a lecturer who changes a weight or a threshold sees every rank update
without rerunning anything.

Like the other engagement modules this one imports neither Streamlit nor pymongo. It does not
read the clock or run git. The caller passes the run time and commit in RunInfo.
"""
from dataclasses import dataclass
from typing import Any, Callable, List, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from utils.engagement import EngagementReport, Measure, PrizeTable, StudentMetrics
from utils.engagement_checks import MAX_DETAIL_LINES, CheckResult

# The prize sheets keep their weights, headings and data on fixed rows. The thresholds sit in
# the two rows under the instructions.
THRESHOLD_ROWS = (4, 5)
WEIGHT_ROW = 7
HEADER_ROW = 8
FIRST_DATA_ROW = 9

# Scores are rounded so that scores equal in exact arithmetic compare as equal in the rank
# formula, whatever order the spreadsheet program adds the terms in.
SCORE_DECIMALS = 6

_EDGE = Side(style="thin", color="BF9000")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_BORDER = Border(left=_EDGE, right=_EDGE, top=_EDGE, bottom=_EDGE)
BOLD = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14)
HEADER_ALIGNMENT = Alignment(wrap_text=True, vertical="center")
GREY_TEXT = Font(color="808080")

TIMING_FLAG = "Tutor timing incomplete"


@dataclass(frozen=True)
class RunInfo:
    """What the caller knows about this run and the About sheet records."""

    run_time: str
    git_commit: str
    checks: Sequence[CheckResult] = ()
    # Results of the checks that need a person, one line each.
    notes: Sequence[str] = ()


@dataclass(frozen=True)
class Column:
    header: str
    read: Callable[[StudentMetrics], Any]
    number_format: Optional[str] = None
    width: float = 13


STUDENT_COLUMNS: Tuple[Column, ...] = (
    Column("Student", lambda s: s.user_id, width=14),
    Column("Tutor timing", lambda s: TIMING_FLAG if s.tutor_timing_incomplete else None, width=14),
    # Overall usage
    Column("Tutor messages", lambda s: s.tutor_messages, "0"),
    Column("Tutor sessions", lambda s: s.tutor_sessions, "0"),
    Column("Tutor active days", lambda s: s.tutor_active_days, "0"),
    Column("Assessor submissions", lambda s: s.assessor_submissions, "0"),
    Column("Assessor sessions", lambda s: s.assessor_sessions, "0"),
    Column("Assessor active days", lambda s: s.assessor_active_days, "0"),
    Column("Sessions, combined", lambda s: s.sessions, "0"),
    Column("Active days, combined", lambda s: s.active_days, "0"),
    # Consistency
    Column("Active weeks, combined", lambda s: s.active_weeks, "0"),
    Column("Active days per week", lambda s: s.active_days_per_week, "0.00"),
    Column("First active", lambda s: s.first_active, "d mmm yyyy"),
    Column("Last active", lambda s: s.last_active, "d mmm yyyy"),
    # Depth
    Column("Tutor interactions per session", lambda s: s.tutor_interactions_per_session, "0.0"),
    Column("Assessor interactions per session", lambda s: s.assessor_interactions_per_session, "0.0"),
    Column("Interactions per session, combined", lambda s: s.interactions_per_session, "0.0"),
    Column("Topic conversations", lambda s: s.topics_worked, "0"),
    Column("Messages per topic conversation", lambda s: s.messages_per_topic, "0.0"),
    Column("Follow-up questions", lambda s: s.follow_ups, "0"),
    Column("Follow-up question share", lambda s: s.follow_up_share, "0%"),
    Column("Assessor questions attempted", lambda s: s.questions_attempted, "0"),
    Column("Questions attempted more than once", lambda s: s.questions_retried, "0"),
    Column("Attempts per question", lambda s: s.attempts_per_question, "0.00"),
    # Meaningful learning engagement
    Column("Questions first scored below 2", lambda s: s.questions_first_below_2, "0"),
    Column("Questions retried after scoring below 2", lambda s: s.questions_responded, "0"),
    Column("Share retried", lambda s: s.responded_share, "0%"),
    Column("Questions improved", lambda s: s.questions_improved, "0"),
    Column("Share improved", lambda s: s.improved_share, "0%"),
    Column("Questions revisited", lambda s: s.questions_revisited, "0"),
    Column("Topics revisited", lambda s: s.topics_revisited, "0"),
    Column("Share of questions and topics revisited", lambda s: s.revisited_share, "0%"),
    Column("Topics at level 2", lambda s: s.topics_at_level_2, "0"),
    Column("Questions at level 2", lambda s: s.questions_at_level_2, "0"),
    # Reconciliation
    Column("Undated tutor messages", lambda s: s.tutor_undated_messages, "0"),
    Column("Tutor error turns", lambda s: s.tutor_error_turns, "0"),
)


# ---------------------------------------------------------------------------
# Cell helpers
# ---------------------------------------------------------------------------

def _put(ws: Worksheet, row: int, column: int, value: Any, number_format: Optional[str] = None):
    """Write a value. None leaves the cell blank, and text is never read as a formula."""
    if value is None:
        return None
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    if number_format:
        cell.number_format = number_format
    return cell


def _formula(ws: Worksheet, row: int, column: int, text: str):
    cell = ws.cell(row=row, column=column)
    cell.value = text
    cell.data_type = "f"
    return cell


def _write_headers(ws: Worksheet, row: int, headers: Sequence[str], height: float = 62) -> None:
    for column, header in enumerate(headers, start=1):
        cell = _put(ws, row, column, header)
        cell.alignment = HEADER_ALIGNMENT
        cell.font = BOLD
    ws.row_dimensions[row].height = height


def _add_table(ws: Worksheet, name: str, first_column: int, last_column: int, header_row: int, last_row: int) -> None:
    """Turn the range into an Excel table with filters. A table needs at least one data row."""
    reference = f"{get_column_letter(first_column)}{header_row}:{get_column_letter(last_column)}{max(last_row, header_row + 1)}"
    if last_row <= header_row:
        ws.auto_filter.ref = reference
        return
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)


def _set_widths(ws: Worksheet, widths: Sequence[float]) -> None:
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width


def _range(first_column: int, last_column: int, first_row: int, last_row: int, absolute: bool = False) -> str:
    mark = "$" if absolute else ""
    return f"{mark}{get_column_letter(first_column)}{mark}{first_row}:{mark}{get_column_letter(last_column)}{mark}{last_row}"


# ---------------------------------------------------------------------------
# Students, Weekly and Attempts
# ---------------------------------------------------------------------------

def _write_students(wb: Workbook, report: EngagementReport) -> None:
    ws = wb.create_sheet("Students")
    _write_headers(ws, 1, [c.header for c in STUDENT_COLUMNS])
    for row, student in enumerate(report.students, start=2):
        for column, spec in enumerate(STUDENT_COLUMNS, start=1):
            _put(ws, row, column, spec.read(student), spec.number_format)
    _set_widths(ws, [c.width for c in STUDENT_COLUMNS])
    _add_table(ws, "Students", 1, len(STUDENT_COLUMNS), 1, len(report.students) + 1)
    ws.freeze_panes = "B2"


def _write_weekly(wb: Workbook, report: EngagementReport) -> None:
    ws = wb.create_sheet("Weekly")
    weeks = report.weeks
    count = len(weeks)
    tutor_first = 2
    undated_column = tutor_first + count
    assessor_first = undated_column + 1
    last_column = assessor_first + count - 1 if count else undated_column

    _put(ws, 1, tutor_first, "Tutor messages in each week, by the Monday that starts the week").font = BOLD
    _put(ws, 1, assessor_first, "Assessor submissions in each week, by the Monday that starts the week").font = BOLD
    _put(ws, 2, 1, "Student").font = BOLD
    for offset, week in enumerate(weeks):
        for first in (tutor_first, assessor_first):
            cell = _put(ws, 2, first + offset, week, "d mmm yy")
            cell.font = BOLD
            cell.alignment = Alignment(horizontal="center")
    undated = _put(ws, 2, undated_column, "Undated")
    undated.font = BOLD
    undated.alignment = Alignment(horizontal="center")

    for row, student in enumerate(report.students, start=3):
        _put(ws, row, 1, student.user_id)
        for offset, week in enumerate(weeks):
            # A week with no activity stays blank so the heat map shows only where students worked.
            _put(ws, row, tutor_first + offset, student.tutor_by_week.get(week) or None, "0")
            _put(ws, row, assessor_first + offset, student.assessor_by_week.get(week) or None, "0")
        _put(ws, row, undated_column, student.tutor_undated_messages or None, "0")

    last_row = len(report.students) + 2
    if count and report.students:
        for first, colour in ((tutor_first, "2F6EB5"), (assessor_first, "2E8B57")):
            ws.conditional_formatting.add(
                _range(first, first + count - 1, 3, last_row),
                ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color=colour),
            )

    ws.column_dimensions["A"].width = 14
    # Wide enough for a date heading and the filter button beside it.
    for column in range(2, last_column + 1):
        ws.column_dimensions[get_column_letter(column)].width = 13
    ws.auto_filter.ref = _range(1, last_column, 2, max(last_row, 3))
    ws.freeze_panes = "B3"


def _write_attempts(wb: Workbook, report: EngagementReport) -> None:
    ws = wb.create_sheet("Attempts")
    headers = [
        "Student", "Module", "Question (first 80 characters)", "Attempts", "Level sequence",
        "First level", "Best level", "Days attempted",
    ]
    _write_headers(ws, 1, headers, height=32)
    for row, attempt in enumerate(report.attempts, start=2):
        _put(ws, row, 1, attempt.user_id)
        _put(ws, row, 2, int(attempt.module) if attempt.module.isdigit() else attempt.module)
        _put(ws, row, 3, attempt.question_start)
        _put(ws, row, 4, attempt.attempts, "0")
        _put(ws, row, 5, attempt.level_sequence)
        _put(ws, row, 6, attempt.first_level, "0")
        _put(ws, row, 7, attempt.best_level, "0")
        _put(ws, row, 8, ", ".join(day.isoformat() for day in attempt.days))
    _set_widths(ws, [14, 9, 70, 10, 16, 10, 10, 40])
    _add_table(ws, "Attempts", 1, len(headers), 1, len(report.attempts) + 1)
    ws.freeze_panes = "B2"


# ---------------------------------------------------------------------------
# Prize sheets
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class _PrizeSpec:
    sheet: str
    table_name: str
    title: str
    description: str
    has_eligibility: bool


PRIZE_1 = _PrizeSpec(
    "Prize 1",
    "Prize1",
    "Prize 1. Most frequent engagement",
    "Every student with at least one interaction appears here. The starting weights rank on "
    "active days and active weeks, which reward coming back over a whole semester.",
    False,
)
PRIZE_2 = _PrizeSpec(
    "Prize 2",
    "Prize2",
    "Prize 2. Most effective use",
    "Only students at or above both thresholds get a rank. Percentiles are worked out across "
    "every student with data, so changing a threshold never needs a rerun.",
    True,
)

PRIZE_INSTRUCTIONS = (
    "Change the yellow cells and every score and rank updates straight away. "
    "A weight of 0 leaves a measure out. Rank 1 is the best."
)


def _raw_headers(measure: Measure) -> List[str]:
    if not measure.is_rate:
        return [f"{measure.label} (raw)"]
    return [
        f"{measure.label} ({measure.successes_label})",
        f"{measure.label} ({measure.opportunities_label})",
        f"{measure.label} (raw rate)",
        f"{measure.label} (adjusted rate)",
    ]


def _rank_formula(score_column: str, row: int, first: int, last: int, eligible_column: Optional[str]) -> str:
    """One plus the number of ranked students with a higher score, which is what RANK gives."""
    scores = f"${score_column}${first}:${score_column}${last}"
    higher = f"ISNUMBER({scores})*({scores}>{score_column}{row})"
    if eligible_column is None:
        return f'=IF(ISNUMBER({score_column}{row}),1+SUMPRODUCT({higher}),"")'
    eligible = f"${eligible_column}${first}:${eligible_column}${last}"
    return (
        f'=IF(AND({eligible_column}{row}="Yes",ISNUMBER({score_column}{row})),'
        f'1+SUMPRODUCT(({eligible}="Yes")*{higher}),"")'
    )


def _score_formula(weights: str, percentiles: str) -> str:
    return (
        f'=IFERROR(ROUND(SUMPRODUCT({weights},{percentiles})'
        f'/SUMPRODUCT({weights},--ISNUMBER({percentiles})),{SCORE_DECIMALS}),"")'
    )


def _write_prize(wb: Workbook, report: EngagementReport, table: PrizeTable, spec: _PrizeSpec) -> None:
    ws = wb.create_sheet(spec.sheet)
    measures = table.measures
    metrics = {student.user_id: student for student in report.students}
    ordered = sorted(table.rows, key=lambda row: (row.rank is None, row.rank or 0, row.user_id))

    # Column layout. The percentile columns are contiguous so the weights row above them lines up
    # one to one, which is what the score formula needs.
    lead = ["Student", "Rank", "Score"] + (["Eligible"] if spec.has_eligibility else []) + ["Tutor timing"]
    rank_column, score_column = 2, 3
    eligible_column = 4 if spec.has_eligibility else None
    percentile_first = len(lead) + 1
    percentile_last = percentile_first + len(measures) - 1
    context = ["Active weeks", "Interactions"] if spec.has_eligibility else []
    context_first = percentile_last + 1
    raw_first = context_first + len(context)

    headers = lead + [f"{m.label} percentile" for m in measures] + context
    raw_start = {}
    for measure in measures:
        raw_start[measure.key] = len(headers) + 1
        headers += _raw_headers(measure)
    last_column = len(headers)
    last_row = FIRST_DATA_ROW + len(ordered) - 1

    _put(ws, 1, 1, spec.title).font = TITLE_FONT
    _put(ws, 2, 1, spec.description)
    _put(ws, 3, 1, PRIZE_INSTRUCTIONS)

    validation_weights = DataValidation(
        type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
        showErrorMessage=True, errorTitle="Weight", error="Enter a weight of 0 or more.",
    )
    validation_count = DataValidation(
        type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
        showErrorMessage=True, errorTitle="Threshold", error="Enter a whole number of 0 or more.",
    )
    ws.add_data_validation(validation_weights)
    ws.add_data_validation(validation_count)

    thresholds = {}
    if spec.has_eligibility:
        settings = report.settings
        for row, label, value in (
            (THRESHOLD_ROWS[0], "Minimum active weeks", settings.min_weeks),
            (THRESHOLD_ROWS[1], "Minimum interactions", settings.min_interactions),
        ):
            _put(ws, row, 1, label).font = BOLD
            cell = _put(ws, row, eligible_column, value, "0")
            cell.fill, cell.border = INPUT_FILL, INPUT_BORDER
            validation_count.add(cell)
            thresholds[label] = f"${get_column_letter(eligible_column)}${row}"

    label_cell = _put(ws, WEIGHT_ROW, percentile_first - 1, "Weights")
    label_cell.font = BOLD
    label_cell.alignment = Alignment(horizontal="right")
    for offset, measure in enumerate(measures):
        cell = _put(ws, WEIGHT_ROW, percentile_first + offset, table.weights[measure.key], "General")
        cell.fill, cell.border = INPUT_FILL, INPUT_BORDER
        validation_weights.add(cell)
    _put(ws, WEIGHT_ROW, raw_first, "Raw values").font = BOLD
    for measure in measures:
        rate = table.cohort_rates.get(measure.key)
        if measure.is_rate and rate is not None:
            _put(ws, WEIGHT_ROW, raw_start[measure.key] + 3, f"Cohort rate {rate:.1%}").font = GREY_TEXT

    _write_headers(ws, HEADER_ROW, headers)

    weights_range = _range(percentile_first, percentile_last, WEIGHT_ROW, WEIGHT_ROW, absolute=True)
    score_letter = get_column_letter(score_column)
    eligible_letter = get_column_letter(eligible_column) if eligible_column else None
    weeks_letter = get_column_letter(context_first)
    interactions_letter = get_column_letter(context_first + 1)

    for row, prize_row in enumerate(ordered, start=FIRST_DATA_ROW):
        student = metrics[prize_row.user_id]
        _put(ws, row, 1, prize_row.user_id)
        _put(ws, row, len(lead), TIMING_FLAG if student.tutor_timing_incomplete else None)

        percentiles = _range(percentile_first, percentile_last, row, row)
        _formula(ws, row, score_column, _score_formula(weights_range, percentiles)).number_format = "0.0"
        _formula(ws, row, rank_column, _rank_formula(score_letter, row, FIRST_DATA_ROW, last_row, eligible_letter)).number_format = "0"
        if spec.has_eligibility:
            _formula(
                ws, row, eligible_column,
                f'=IF(AND({weeks_letter}{row}>={thresholds["Minimum active weeks"]},'
                f'{interactions_letter}{row}>={thresholds["Minimum interactions"]}),"Yes","No")',
            )
            _put(ws, row, context_first, student.active_weeks, "0")
            _put(ws, row, context_first + 1, student.interactions, "0")

        for offset, measure in enumerate(measures):
            cell = prize_row.cells[measure.key]
            _put(ws, row, percentile_first + offset, cell.percentile, "0.0")
            start = raw_start[measure.key]
            if measure.is_rate:
                _put(ws, row, start, cell.successes, "0")
                _put(ws, row, start + 1, cell.opportunities, "0")
                _put(ws, row, start + 2, cell.raw, "0.0%")
                _put(ws, row, start + 3, cell.adjusted, "0.0%")
            else:
                _put(ws, row, start, cell.raw, "0" if isinstance(cell.raw, int) else "0.00")

    if ordered and spec.has_eligibility:
        # Students below a threshold stay on the sheet but are greyed out and have no rank.
        ws.conditional_formatting.add(
            _range(1, last_column, FIRST_DATA_ROW, last_row),
            FormulaRule(formula=[f'${eligible_letter}{FIRST_DATA_ROW}="No"'], font=GREY_TEXT),
        )

    widths = [14, 8, 9] + ([10] if spec.has_eligibility else []) + [22]
    widths += [13] * len(measures) + [11] * len(context) + [14] * (last_column - raw_first + 1)
    _set_widths(ws, widths)
    _add_table(ws, spec.table_name, 1, last_column, HEADER_ROW, last_row)
    ws.freeze_panes = f"B{FIRST_DATA_ROW}"


# ---------------------------------------------------------------------------
# About
# ---------------------------------------------------------------------------

INTRO = (
    "This workbook gives the PRQ lecturers student-level data for two prizes. One is for the "
    "students who used the AI tutor most often. The other is for the students who used it most "
    "effectively to support their learning. The AI tutor means the whole app, so both the Tutor "
    "and the Assessor count. The workbook sets out the data. The lecturers decide how to weight "
    "it and who wins."
)

SHEETS = (
    ("Students", "One row per student with every usage, consistency, depth and learning measure. "
                 "Tutor and Assessor figures sit side by side and are never added into one count."),
    ("Weekly", "Interactions in each week from the first week of activity. The Tutor and Assessor "
               "blocks are shaded as heat maps. Tutor messages logged before timestamps were added "
               "have no date, so they sit in the Undated column."),
    ("Attempts", "One row per student and assessor question with the attempt count, the sequence of "
                 "levels, the first and best level and the days attempted. Every learning figure on "
                 "the Students sheet can be traced to these rows."),
    ("Prize 1", "The measures for most frequent engagement, with a weight for each."),
    ("Prize 2", "The measures for most effective use, with a weight for each and thresholds for who "
                "can be ranked."),
)

CHANGING_WEIGHTS = (
    "On the Prize 1 and Prize 2 sheets, the yellow cells in the Weights row set how much each "
    "measure counts. Type a new weight and every score and rank updates straight away. A weight "
    "of 0 leaves the measure out.",
    "The score is the weighted mean of the student's percentiles. A blank percentile means the "
    "student has no data for that measure, and the mean leaves it out instead of counting it as "
    "zero. A student who never used the Tutor isn't marked down on tutor measures.",
    "Prize 2 has two more yellow cells above the weights, for the minimum active weeks and the "
    "minimum interactions. A student below either one keeps their row, greyed out, with no rank.",
    "Scores and ranks are formulas, so they appear once the file is opened in Excel or LibreOffice. "
    "A viewer that doesn't calculate formulas shows them blank.",
    "The percentiles, the shrinkage constant and the session gap are worked out by the script. "
    "Changing those means running the script again with a different setting.",
    "The starting weights are a suggestion so the sheets rank as soon as they open. Prize 1 ranks "
    "on active days and active weeks because those show how often a student came back, while "
    "interaction counts reward one long sitting. Prize 2 starts with improvement after feedback, "
    "revisiting and follow-up questions because they best separate purposeful use from frequent "
    "use. Level 2 counts start at 0 because a high level can reflect what a student knew before "
    "using the app.",
)


def _definitions(report: EngagementReport) -> Tuple[Tuple[str, str], ...]:
    settings = report.settings
    minutes = settings.session_gap.total_seconds() / 60
    return (
        ("Time", f"Every timestamp is converted to {settings.timezone} time before days and weeks "
                 "are worked out. Weeks run from Monday to Sunday and are labelled by their Monday."),
        ("Interaction", "For the Tutor, a message the student wrote. Opening questions from the tutor "
                        "don't count, and neither does a turn that ended in an error, because the "
                        "student's retry is counted instead. For the Assessor, one graded submission. "
                        "A tutor message can be a single word while a submission is a worked answer, so "
                        "the two are not comparable. Read them side by side and never rank on their sum."),
        ("Session", f"A run of dated interactions with no gap longer than {minutes:g} minutes. Sessions "
                    "are counted for each mode and for both together. Tutor sessions only exist for "
                    "messages logged after timestamps were added."),
        ("Active day and active week", "A date or week with at least one interaction. Tutor messages "
                                       "logged before timestamps were added give two lower-bound dates "
                                       "for each topic, the day the conversation started and the day "
                                       "the topic's level last changed. A student with "
                                       f"{settings.incomplete_timing_messages} or more undated tutor "
                                       "messages is flagged as tutor timing incomplete, because their "
                                       "tutor days are undercounted."),
        ("Active days per week", "Combined active days divided by the number of whole Monday to Sunday "
                                 "weeks from the student's first activity to the report date, counting "
                                 "both ends. A student who first appears in the report week has one week."),
        ("Question", "An assessor question is identified by its module and its text. Each grade is "
                     "converted back to the 0, 1, 2 competency scale."),
        ("Follow-up question", "A tutor message that contains a question mark or starts with a question "
                               "word. This is a word-list heuristic and some flagged messages are not "
                               "real questions."),
        ("Responded to feedback", "A question whose first attempt scored below 2 and that the student "
                                  "tried again."),
        ("Improved through attempts", "A question the student responded to where a later attempt "
                                      "scored higher than the first."),
        ("Revisited", "A question or tutor topic the student worked on across two or more days. A "
                      "return to a question already at level 2 counts, because that is spaced practice."),
        ("Percentile", "The student's position in the cohort from 0 to 100, with tied students sharing "
                       "the average rank. The lowest value gets 0 and the highest gets 100. A student "
                       "alone in the cohort, or a cohort where every value ties, gets 50."),
        ("Shrinkage", f"A rate such as improvement after feedback is shrunk toward the cohort rate "
                      f"before its percentile is worked out, so one lucky attempt can't reach the top. "
                      f"The adjusted rate is (successes + k × cohort rate) / (opportunities + k), with "
                      f"k set to {settings.shrink_k:g} for this run."),
        ("Rank", "Students with the same score share the higher rank, so scores of 90, 80, 80 and 70 "
                 "rank 1, 2, 2 and 4."),
    )


KNOWN_LIMITS = (
    "Tutor days and revisits from before timestamps were added are lower bounds. Look at flagged "
    "students by hand.",
    "Follow-up questions are found with a word list.",
    "Changing a question's text after students have attempted it splits its attempts across two rows.",
    "Prize 2 ranks a student who only used the Tutor on weaker measures, because they rest on a word "
    "list and lower-bound dates. Treat their rank with more caution.",
    "Staff and test accounts appear in the output. Read the Attempts rows and a sample of tutor "
    "conversations for the students at the top of each sheet before deciding.",
)


class _AboutWriter:
    def __init__(self, ws: Worksheet):
        self.ws = ws
        self.row = 1
        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 110

    def title(self, text: str) -> None:
        _put(self.ws, self.row, 1, text).font = TITLE_FONT
        self.row += 2

    def section(self, text: str) -> None:
        self.row += 1
        _put(self.ws, self.row, 1, text).font = Font(bold=True, size=12)
        self.row += 1

    def paragraph(self, text: str) -> None:
        cell = _put(self.ws, self.row, 1, text)
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=2)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        # Merged cells don't auto-fit, so the height is set from the length of the text.
        self.ws.row_dimensions[self.row].height = 15 * (len(text) // 135 + 1)
        self.row += 1

    def pair(self, label: str, value: Any, number_format: Optional[str] = None) -> None:
        label_cell = _put(self.ws, self.row, 1, label)
        if label_cell is not None:
            label_cell.font = BOLD
            label_cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell = _put(self.ws, self.row, 2, value, number_format)
        if cell is not None:
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        self.row += 1


def _write_about(wb: Workbook, report: EngagementReport, run: RunInfo) -> None:
    ws = wb.active
    ws.title = "About"
    about = _AboutWriter(ws)
    settings = report.settings

    about.title("PRQ AI tutor engagement report")
    about.paragraph(INTRO)

    about.section("This run")
    about.pair("Run time", run.run_time)
    about.pair("Report date", report.report_date, "d mmm yyyy")
    if report.data_range:
        first, last = report.data_range
        about.pair("Data range", f"{first:%d %b %Y} to {last:%d %b %Y}")
    else:
        about.pair("Data range", "No activity was found.")
    about.pair("Students in the report", len(report.students))
    about.pair("Git commit", run.git_commit)
    about.pair("Session gap, minutes", settings.session_gap.total_seconds() / 60)
    about.pair("Shrinkage constant k", settings.shrink_k)
    about.pair("Minimum active weeks, starting value", settings.min_weeks)
    about.pair("Minimum interactions, starting value", settings.min_interactions)
    about.pair("Undated tutor messages that set the timing flag", settings.incomplete_timing_messages)
    about.pair("Time zone for days and weeks", settings.timezone)
    about.pair("Time zone of progress records", settings.progress_timezone)

    about.section("Sheets")
    for name, text in SHEETS:
        about.pair(name, text)

    about.section("Changing weights and thresholds")
    for text in CHANGING_WEIGHTS:
        about.paragraph(text)

    about.section("Definitions")
    for name, text in _definitions(report):
        about.pair(name, text)

    about.section("Known limits")
    for text in KNOWN_LIMITS:
        about.paragraph(text)

    about.section("Validation")
    if not run.checks and not run.notes:
        about.paragraph("No validation checks were recorded for this run.")
    for check in run.checks:
        about.pair(f"{check.name}. {check.status}", check.summary)
        for line in check.details[: MAX_DETAIL_LINES + 1]:
            about.pair("", line)
    for note in run.notes:
        about.pair("Checked by hand", note)

    if report.warnings:
        about.section("Warnings")
        for line in report.warnings[:50]:
            about.pair("", line)
        if len(report.warnings) > 50:
            about.pair("", f"... and {len(report.warnings) - 50} more")


def build_workbook(report: EngagementReport, run: RunInfo) -> Workbook:
    """Build the whole workbook. The caller saves it."""
    wb = Workbook()
    _write_about(wb, report, run)
    _write_students(wb, report)
    _write_weekly(wb, report)
    _write_attempts(wb, report)
    _write_prize(wb, report, report.prize_1, PRIZE_1)
    _write_prize(wb, report, report.prize_2, PRIZE_2)
    wb.properties.creator = "PRQ engagement report"
    wb.properties.title = "PRQ AI tutor engagement report"
    return wb
